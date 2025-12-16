#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador de Presentación PowerPoint sobre Operaciones de Rescate
Universidad Fidelitas - Curso de Fundamentos de Sistemas Operativos
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def crear_presentacion_rescate():
    """Crea una presentación PowerPoint completa sobre operaciones de rescate"""
    
    # Crear presentación
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Colores corporativos
    COLOR_TITULO = RGBColor(192, 0, 0)  # Rojo rescate
    COLOR_SUBTITULO = RGBColor(51, 51, 51)  # Gris oscuro
    COLOR_TEXTO = RGBColor(64, 64, 64)  # Gris medio
    COLOR_ACENTO = RGBColor(255, 140, 0)  # Naranja
    
    # ========== DIAPOSITIVA 1: PORTADA ==========
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Layout en blanco
    
    # Fondo de color
    background = slide1.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(240, 240, 240)
    
    # Título principal
    titulo_box = slide1.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1))
    titulo_frame = titulo_box.text_frame
    titulo_frame.text = "OPERACIONES DE RESCATE"
    titulo_p = titulo_frame.paragraphs[0]
    titulo_p.font.size = Pt(54)
    titulo_p.font.bold = True
    titulo_p.font.color.rgb = COLOR_TITULO
    titulo_p.alignment = PP_ALIGN.CENTER
    
    # Subtítulo
    subtitulo_box = slide1.shapes.add_textbox(Inches(1), Inches(3.8), Inches(8), Inches(0.6))
    subtitulo_frame = subtitulo_box.text_frame
    subtitulo_frame.text = "Protocolos, Técnicas y Procedimientos"
    subtitulo_p = subtitulo_frame.paragraphs[0]
    subtitulo_p.font.size = Pt(28)
    subtitulo_p.font.color.rgb = COLOR_SUBTITULO
    subtitulo_p.alignment = PP_ALIGN.CENTER
    
    # Información del curso
    info_box = slide1.shapes.add_textbox(Inches(1), Inches(5.5), Inches(8), Inches(1))
    info_frame = info_box.text_frame
    info_frame.text = "Universidad Fidelitas\nCurso de Fundamentos de Sistemas Operativos\n2025"
    for paragraph in info_frame.paragraphs:
        paragraph.font.size = Pt(18)
        paragraph.font.color.rgb = COLOR_TEXTO
        paragraph.alignment = PP_ALIGN.CENTER
    
    # ========== DIAPOSITIVA 2: ÍNDICE ==========
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    
    titulo2 = slide2.shapes.title
    titulo2.text = "Contenido"
    titulo2.text_frame.paragraphs[0].font.size = Pt(44)
    titulo2.text_frame.paragraphs[0].font.color.rgb = COLOR_TITULO
    
    contenido2 = slide2.placeholders[1]
    tf2 = contenido2.text_frame
    tf2.clear()
    
    items_indice = [
        "1. Tipos de Operaciones de Rescate",
        "2. Protocolos de Seguridad",
        "3. Equipamiento Esencial",
        "4. Fases de una Operación",
        "5. Coordinación y Comunicación",
        "6. Conclusiones"
    ]
    
    for item in items_indice:
        p = tf2.add_paragraph()
        p.text = item
        p.font.size = Pt(24)
        p.font.color.rgb = COLOR_TEXTO
        p.space_before = Pt(12)
        p.level = 0
    
    # ========== DIAPOSITIVA 3: TIPOS DE RESCATE ==========
    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    
    titulo3 = slide3.shapes.title
    titulo3.text = "Tipos de Operaciones de Rescate"
    titulo3.text_frame.paragraphs[0].font.size = Pt(40)
    titulo3.text_frame.paragraphs[0].font.color.rgb = COLOR_TITULO
    
    contenido3 = slide3.placeholders[1]
    tf3 = contenido3.text_frame
    tf3.clear()
    
    tipos_rescate = [
        ("🏔️ Rescate en Montaña", "Operaciones en terreno elevado y difícil acceso"),
        ("🌊 Rescate Acuático", "Salvamento en ríos, lagos, mares y zonas inundadas"),
        ("🏙️ Rescate Urbano", "Emergencias en edificios, estructuras colapsadas"),
        ("🚁 Rescate Aéreo", "Evacuación y transporte mediante helicópteros"),
        ("🔥 Rescate en Incendios", "Extracción de víctimas en estructuras en llamas"),
        ("⚠️ Rescate Industrial", "Espacios confinados y ambientes peligrosos")
    ]
    
    for tipo, descripcion in tipos_rescate:
        p = tf3.add_paragraph()
        p.text = tipo
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACENTO
        p.space_before = Pt(10)
        
        p2 = tf3.add_paragraph()
        p2.text = descripcion
        p2.font.size = Pt(18)
        p2.font.color.rgb = COLOR_TEXTO
        p2.level = 1
    
    # ========== DIAPOSITIVA 4: PROTOCOLOS DE SEGURIDAD ==========
    slide4 = prs.slides.add_slide(prs.slide_layouts[1])
    
    titulo4 = slide4.shapes.title
    titulo4.text = "Protocolos de Seguridad"
    titulo4.text_frame.paragraphs[0].font.size = Pt(40)
    titulo4.text_frame.paragraphs[0].font.color.rgb = COLOR_TITULO
    
    contenido4 = slide4.placeholders[1]
    tf4 = contenido4.text_frame
    tf4.clear()
    
    protocolos = [
        "✓ Evaluación inicial de la escena",
        "✓ Identificación de peligros potenciales",
        "✓ Establecimiento de perímetro de seguridad",
        "✓ Uso obligatorio de equipo de protección personal (EPP)",
        "✓ Comunicación constante con el equipo",
        "✓ Verificación de condiciones ambientales",
        "✓ Plan de evacuación de emergencia",
        "✓ Monitoreo continuo de la situación"
    ]
    
    for protocolo in protocolos:
        p = tf4.add_paragraph()
        p.text = protocolo
        p.font.size = Pt(22)
        p.font.color.rgb = COLOR_TEXTO
        p.space_before = Pt(8)
    
    # ========== DIAPOSITIVA 5: EQUIPAMIENTO ==========
    slide5 = prs.slides.add_slide(prs.slide_layouts[1])
    
    titulo5 = slide5.shapes.title
    titulo5.text = "Equipamiento Esencial"
    titulo5.text_frame.paragraphs[0].font.size = Pt(40)
    titulo5.text_frame.paragraphs[0].font.color.rgb = COLOR_TITULO
    
    contenido5 = slide5.placeholders[1]
    tf5 = contenido5.text_frame
    tf5.clear()
    
    equipamiento = [
        ("Protección Personal", "Cascos, guantes, arneses, botas especializadas"),
        ("Comunicación", "Radios, GPS, señales de emergencia"),
        ("Herramientas de Corte", "Sierras, cortadores hidráulicos, esparcidores"),
        ("Equipo de Elevación", "Cuerdas, poleas, camillas, sistemas de anclaje"),
        ("Iluminación", "Linternas, reflectores, señales luminosas"),
        ("Primeros Auxilios", "Botiquín completo, desfibrilador, oxígeno")
    ]
    
    for categoria, items in equipamiento:
        p = tf5.add_paragraph()
        p.text = categoria
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACENTO
        p.space_before = Pt(10)
        
        p2 = tf5.add_paragraph()
        p2.text = items
        p2.font.size = Pt(18)
        p2.font.color.rgb = COLOR_TEXTO
        p2.level = 1
    
    # ========== DIAPOSITIVA 6: FASES DE OPERACIÓN ==========
    slide6 = prs.slides.add_slide(prs.slide_layouts[1])
    
    titulo6 = slide6.shapes.title
    titulo6.text = "Fases de una Operación de Rescate"
    titulo6.text_frame.paragraphs[0].font.size = Pt(40)
    titulo6.text_frame.paragraphs[0].font.color.rgb = COLOR_TITULO
    
    contenido6 = slide6.placeholders[1]
    tf6 = contenido6.text_frame
    tf6.clear()
    
    fases = [
        ("1️⃣ ACTIVACIÓN", "Recepción de alerta y movilización del equipo"),
        ("2️⃣ EVALUACIÓN", "Análisis de la situación y planificación táctica"),
        ("3️⃣ PREPARACIÓN", "Organización de recursos y equipamiento"),
        ("4️⃣ ACCESO", "Llegada al lugar y establecimiento de zona segura"),
        ("5️⃣ ESTABILIZACIÓN", "Control de riesgos y aseguramiento del área"),
        ("6️⃣ RESCATE", "Extracción de víctimas de forma segura"),
        ("7️⃣ EVACUACIÓN", "Traslado a zona segura o centro médico"),
        ("8️⃣ CIERRE", "Evaluación post-operación y reporte")
    ]
    
    for fase, descripcion in fases:
        p = tf6.add_paragraph()
        p.text = fase
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACENTO
        p.space_before = Pt(6)
        
        p2 = tf6.add_paragraph()
        p2.text = descripcion
        p2.font.size = Pt(16)
        p2.font.color.rgb = COLOR_TEXTO
        p2.level = 1
    
    # ========== DIAPOSITIVA 7: COORDINACIÓN ==========
    slide7 = prs.slides.add_slide(prs.slide_layouts[1])
    
    titulo7 = slide7.shapes.title
    titulo7.text = "Coordinación y Comunicación"
    titulo7.text_frame.paragraphs[0].font.size = Pt(40)
    titulo7.text_frame.paragraphs[0].font.color.rgb = COLOR_TITULO
    
    contenido7 = slide7.placeholders[1]
    tf7 = contenido7.text_frame
    tf7.clear()
    
    coordinacion = [
        "📡 Sistema de Comando de Incidentes (SCI)",
        "🎯 Designación clara de roles y responsabilidades",
        "📞 Canales de comunicación establecidos",
        "🤝 Coordinación interinstitucional",
        "📋 Registro y documentación de acciones",
        "⏱️ Actualizaciones periódicas del estado",
        "🚨 Protocolos de escalamiento",
        "👥 Trabajo en equipo y apoyo mutuo"
    ]
    
    for item in coordinacion:
        p = tf7.add_paragraph()
        p.text = item
        p.font.size = Pt(22)
        p.font.color.rgb = COLOR_TEXTO
        p.space_before = Pt(10)
    
    # ========== DIAPOSITIVA 8: CONCLUSIONES ==========
    slide8 = prs.slides.add_slide(prs.slide_layouts[1])
    
    titulo8 = slide8.shapes.title
    titulo8.text = "Conclusiones"
    titulo8.text_frame.paragraphs[0].font.size = Pt(40)
    titulo8.text_frame.paragraphs[0].font.color.rgb = COLOR_TITULO
    
    contenido8 = slide8.placeholders[1]
    tf8 = contenido8.text_frame
    tf8.clear()
    
    conclusiones = [
        "Las operaciones de rescate requieren preparación exhaustiva y entrenamiento continuo",
        "La seguridad del equipo de rescate es prioritaria en toda operación",
        "La coordinación efectiva salva vidas y optimiza recursos",
        "El equipamiento adecuado es fundamental para el éxito",
        "La comunicación clara previene errores y accidentes",
        "Cada operación debe ser evaluada para mejorar procedimientos futuros"
    ]
    
    for conclusion in conclusiones:
        p = tf8.add_paragraph()
        p.text = "• " + conclusion
        p.font.size = Pt(20)
        p.font.color.rgb = COLOR_TEXTO
        p.space_before = Pt(12)
    
    # ========== DIAPOSITIVA 9: CIERRE ==========
    slide9 = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Fondo
    background9 = slide9.background
    fill9 = background9.fill
    fill9.solid()
    fill9.fore_color.rgb = RGBColor(240, 240, 240)
    
    # Mensaje final
    final_box = slide9.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(2))
    final_frame = final_box.text_frame
    final_frame.text = "¡GRACIAS!"
    final_p = final_frame.paragraphs[0]
    final_p.font.size = Pt(60)
    final_p.font.bold = True
    final_p.font.color.rgb = COLOR_TITULO
    final_p.alignment = PP_ALIGN.CENTER
    
    # Contacto
    contacto_box = slide9.shapes.add_textbox(Inches(1), Inches(5), Inches(8), Inches(1))
    contacto_frame = contacto_box.text_frame
    contacto_frame.text = "Universidad Fidelitas\nCurso de Fundamentos de Sistemas Operativos"
    for paragraph in contacto_frame.paragraphs:
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = COLOR_TEXTO
        paragraph.alignment = PP_ALIGN.CENTER
    
    # Guardar presentación
    nombre_archivo = "Presentacion_Rescate.pptx"
    prs.save(nombre_archivo)
    print(f"✅ Presentación creada exitosamente: {nombre_archivo}")
    print(f"📊 Total de diapositivas: {len(prs.slides)}")
    return nombre_archivo

if __name__ == "__main__":
    crear_presentacion_rescate()
